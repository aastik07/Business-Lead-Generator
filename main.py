# -*- coding: utf-8 -*-
"""
Local Business Lead Generator
-------------------------------
Collects publicly displayed business information from Google Maps
and exports it to a clean Excel file.

Usage:
    python main.py
"""

import os
import re
import time
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import WebDriverException, TimeoutException


# ================================================
#  HELPER — Extract stable Place ID from a Maps URL
# ================================================

def extract_place_id(url):
    """
    Pull the unique Place ID out of a Google Maps URL.

    Every Google Maps listing URL contains a stable identifier in the form:
        !1s0x<hex>:<hex>
    e.g. !1s0x3916d9ccf02a2717:0xd818554986f761d3

    This ID never changes regardless of viewport coordinates (@lat,lng,zoom),
    so it is the most reliable deduplication key.
    Returns the ID string, or None if not found.
    """
    match = re.search(r"!1s(0x[0-9a-fA-F]+:[0-9a-fA-Fx]+)", url)
    return match.group(1) if match else None


# ================================================
#  1. USER INPUT
# ================================================

def get_user_input():
    """Ask the user for city, business niche, and max leads. Validate each."""

    print("\n========================================")
    print("   LOCAL BUSINESS LEAD GENERATOR")
    print("========================================\n")

    # City
    while True:
        city = input("Enter city: ").strip()
        if city:
            break
        print("  [!] City cannot be empty. Please try again.\n")

    # Business niche
    while True:
        niche = input("Enter business niche: ").strip()
        if niche:
            break
        print("  [!] Business niche cannot be empty. Please try again.\n")

    # Maximum leads
    while True:
        raw = input("Enter maximum leads: ").strip()
        if raw.isdigit() and int(raw) > 0:
            max_leads = int(raw)
            break
        print("  [!] Please enter a valid positive number.\n")

    return city, niche, max_leads


# Chrome Setup

def start_driver():
    """Start Chrome with Selenium. Returns the driver, or None on failure."""

    print("\nStarting Chrome...")

    try:
        options = Options()
        options.add_argument("--start-maximized")
        # Suppress unnecessary console noise from Chrome
        options.add_experimental_option("excludeSwitches", ["enable-logging"])
        driver = webdriver.Chrome(options=options)
        return driver
    except WebDriverException as e:
        print("\n[ERROR] Could not start Chrome.")
        print("  Make sure Google Chrome and ChromeDriver are installed.")
        print(f"  Details: {e}")
        return None


# Google Map Search

def open_google_maps(driver, city, niche):
    """
    Navigate to Google Maps with the search query.
    Returns the scrollable results panel, or None if not found.
    """

    search_query = f"{niche} in {city}"
    maps_url = (
        "https://www.google.com/maps/search/"
        + niche.replace(" ", "+")
        + "+in+"
        + city.replace(" ", "+")
    )

    print(f"\nSearching Google Maps for:\n{search_query}\n")

    try:
        driver.get(maps_url)
        time.sleep(5)  # Wait for the page to load
    except WebDriverException:
        print("[ERROR] Could not load Google Maps. Check your internet connection.")
        return None

    # Check for CAPTCHA or consent blocking page
    current = driver.current_url.lower()
    if "consent.google" in current or "sorry" in current:
        print("[ERROR] Google is showing a consent or CAPTCHA page.")
        print("  Please open Chrome manually, accept cookies, then re-run.")
        return None

    # Find the scrollable results panel
    try:
        panel = driver.find_element(
            By.XPATH, '//div[contains(@aria-label, "Results for")]'
        )
        return panel
    except Exception:
        print("[ERROR] No results panel found.")
        print("  Google Maps may have changed its layout,")
        print("  or there are no results for this search.")
        return None


# Extracting Business Details

def extract_business_info(driver, item, city, niche):
    """
    Click a listing card and extract all available business details.
    Uses WebDriverWait instead of fixed sleep for speed.
    Returns a dict, using 'N/A' for any field that cannot be found.
    """

    info = {
        "Business Name":  "N/A",
        "Rating":         "N/A",
        "Reviews":        "N/A",
        "Phone":          "N/A",
        "Address":        "N/A",
        "Plus Code":      "N/A",
        "City":           city,
        "Business Niche": niche,
        "Google Maps URL": "N/A",
    }

    try:
        # Scroll the card into view and click it
        driver.execute_script("arguments[0].scrollIntoView(true);", item)
        time.sleep(0.3)
        item.click()

        # Wait until the business name heading appears (max 6 seconds)
        # This replaces the old fixed time.sleep(3) — only waits as long as needed
        try:
            WebDriverWait(driver, 6).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "h1.DUwDvf"))
            )
        except TimeoutException:
            # Detail panel did not load in time — skip this listing
            return info

        # Business name
        try:
            info["Business Name"] = driver.find_element(
                By.CSS_SELECTOR, "h1.DUwDvf"
            ).text.strip()
        except Exception:
            pass

        # Rating (e.g. "4.5")
        try:
            info["Rating"] = driver.find_element(
                By.CSS_SELECTOR, "div.F7nice span[aria-hidden='true']"
            ).text.strip()
        except Exception:
            pass

        # Review count — try two selectors; Google Maps occasionally changes class names
        # Method 1: span with aria-label like "1,234 reviews" inside the rating block
        try:
            reviews_el = driver.find_element(
                By.CSS_SELECTOR, "div.F7nice span[aria-label]"
            )
            aria = reviews_el.get_attribute("aria-label") or ""
            # aria-label looks like "1,234 reviews" — grab the number part
            info["Reviews"] = aria.replace(" reviews", "").replace(",", "").strip()
        except Exception:
            pass

        # Method 2 fallback: UY7F9 span (review count in parentheses like "(312)")
        if info["Reviews"] == "N/A":
            try:
                raw = driver.find_element(By.CSS_SELECTOR, "span.UY7F9").text.strip()
                info["Reviews"] = raw.strip("()")
            except Exception:
                pass

        # Phone number
        try:
            phone_btn = driver.find_element(
                By.CSS_SELECTOR, "button[data-tooltip='Copy phone number']"
            )
            raw = phone_btn.get_attribute("aria-label") or ""
            info["Phone"] = raw.replace("Phone: ", "").strip()
        except Exception:
            pass

        # Address
        try:
            addr_btn = driver.find_element(
                By.CSS_SELECTOR, "button[data-tooltip='Copy address']"
            )
            raw = addr_btn.get_attribute("aria-label") or ""
            info["Address"] = raw.replace("Address: ", "").strip()
        except Exception:
            pass

        # Plus Code
        try:
            plus_btn = driver.find_element(
                By.CSS_SELECTOR, "button[data-tooltip='Copy plus code']"
            )
            raw = plus_btn.get_attribute("aria-label") or ""
            info["Plus Code"] = raw.replace("Plus code: ", "").strip()
        except Exception:
            pass

        # Google Maps URL
        # Wait briefly so Chrome finishes updating the address bar to THIS listing.
        # Without this pause, driver.current_url can still show the previous listing's URL.
        try:
            time.sleep(0.8)
            info["Google Maps URL"] = driver.current_url
        except Exception:
            pass

    except Exception:
        # If the whole item fails, return what we have (mostly N/A)
        pass

    return info


# Collecting & Scrolling

def scroll_and_collect(driver, panel, city, niche, max_leads):
    """
    Scroll through the Google Maps results panel and collect business leads.
    Stops when max_leads is reached or no new results appear after retries.
    Returns a list of business info dicts.
    """

    print("Collecting businesses...\n")

    leads = []
    seen_place_ids  = set()   # Primary:   Google Place ID (stable, from URL)
    seen_name_phone = set()   # Secondary: Name + Phone (unique per real business)
    seen_name_addr  = set()   # Tertiary:  Name + Address (original fallback)

    processed_count = 0    # How many listing cards we have already clicked
    no_new_attempts = 0    # Consecutive scrolls with no new cards
    MAX_ATTEMPTS = 5       # Stop after this many empty scrolls in a row

    while len(leads) < max_leads and no_new_attempts < MAX_ATTEMPTS:

        # Get all currently visible listing cards
        items = driver.find_elements(By.CSS_SELECTOR, "div.Nv2PK")
        new_items = items[processed_count:]

        if not new_items:
            # Scroll down and try again
            no_new_attempts += 1
            driver.execute_script(
                "arguments[0].scrollTop = arguments[0].scrollHeight", panel
            )
            time.sleep(2)
            continue

        # New cards appeared, reset the empty-scroll counter
        no_new_attempts = 0

        for item in new_items:

            if len(leads) >= max_leads:
                break

            info = extract_business_info(driver, item, city, niche)

            # Skip if we could not get even the business name
            if info["Business Name"] == "N/A":
                continue

            # ------ 3-tier duplicate check ----------------------------
            url         = info["Google Maps URL"]
            place_id    = extract_place_id(url) if url != "N/A" else None
            name_phone  = f"{info['Business Name']}|{info['Phone']}"
            name_addr   = f"{info['Business Name']}|{info['Address']}"

            # Tier 1: Place ID — most reliable, immune to viewport changes
            if place_id:
                if place_id in seen_place_ids:
                    print("  Duplicate skipped.")
                    continue
                seen_place_ids.add(place_id)

            # Tier 2: Name + Phone — catches duplicates when URL had no Place ID
            elif info["Phone"] != "N/A":
                if name_phone in seen_name_phone:
                    print("  Duplicate skipped.")
                    continue
                seen_name_phone.add(name_phone)

            # Tier 3: Name + Address — final fallback
            else:
                if name_addr in seen_name_addr:
                    print("  Duplicate skipped.")
                    continue
                seen_name_addr.add(name_addr)
            # -----------------------------------------------------------

            leads.append(info)
            print(f"  [{len(leads)}] {info['Business Name']}")

        # Mark all cards seen so far as processed
        processed_count = len(items)

        # Scroll down to load more results
        driver.execute_script(
            "arguments[0].scrollTop = arguments[0].scrollHeight", panel
        )
        time.sleep(2)

    return leads


# Saving Excel Files to Output Folder

def save_to_excel(leads, city, niche):
    """
    Save the collected leads to a styled Excel file inside the output/ folder.
    Returns the file path on success, or None on failure.
    """

    # Create the output/ folder if it does not exist yet
    output_folder = "output"
    os.makedirs(output_folder, exist_ok=True)

    # Build filename: interior_designers_chandigarh_2026-09-08.xlsx
    today      = datetime.today().strftime("%Y-%m-%d")
    safe_niche = niche.lower().replace(" ", "_")
    safe_city  = city.lower().replace(" ", "_")
    filename   = f"{safe_niche}_{safe_city}_{today}.xlsx"
    filepath   = os.path.join(output_folder, filename)

    # Column order matches the dict keys
    headers = [
        "Business Name", "Rating", "Reviews", "Phone", "Address",
        "Plus Code", "City", "Business Niche", "Google Maps URL",
    ]

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"

        # Style the header row
        header_font  = Font(bold=True, color="FFFFFF")
        header_fill  = PatternFill(fill_type="solid", fgColor="1F4E79")  # Dark blue
        header_align = Alignment(horizontal="center", vertical="center")

        for col_index, header in enumerate(headers, start=1):
            cell            = ws.cell(row=1, column=col_index, value=header)
            cell.font       = header_font
            cell.fill       = header_fill
            cell.alignment  = header_align

        # Write data rows
        url_col_index = headers.index("Google Maps URL") + 1  # 1-based column number

        for row_index, lead in enumerate(leads, start=2):
            for col_index, header in enumerate(headers, start=1):
                value = lead[header]
                cell  = ws.cell(row=row_index, column=col_index, value=value)

                # Make the Google Maps URL column a clickable hyperlink
                if col_index == url_col_index and value != "N/A":
                    cell.hyperlink = value
                    cell.font = Font(color="0563C1", underline="single")

        # Freeze the first (header) row
        ws.freeze_panes = "A2"

        # Enable auto-filter on all columns
        ws.auto_filter.ref = ws.dimensions

        # Auto-adjust column widths based on content
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    cell_len = len(str(cell.value)) if cell.value else 0
                    if cell_len > max_length:
                        max_length = cell_len
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

        wb.save(filepath)
        return filepath

    except Exception as e:
        print(f"\n[ERROR] Could not save Excel file: {e}")
        return None


# Main

def main():
    # Step 1: Get validated input from the user
    city, niche, max_leads = get_user_input()

    # Step 2: Start Chrome
    driver = start_driver()
    if driver is None:
        return  # Chrome failed to start; message already shown

    try:
        # Step 3: Open Google Maps and find the results panel
        panel = open_google_maps(driver, city, niche)
        if panel is None:
            return  # Something went wrong; message already shown

        # Step 4: Scroll through results and collect leads
        leads = scroll_and_collect(driver, panel, city, niche, max_leads)

        # Step 5: Report and save
        if not leads:
            print("\nNo leads were collected. Nothing to save.")
            return

        print(f"\nCompleted!")
        print(f"Total leads collected: {len(leads)}\n")

        filepath = save_to_excel(leads, city, niche)

        if filepath:
            print(f"Excel saved to:\n{filepath}\n")

    except KeyboardInterrupt:
        print("\n\nStopped by user.")

    finally:
        # Always close Chrome, even if an error occurred
        try:
            driver.quit()
        except Exception:
            pass


if __name__ == "__main__":
    main()
